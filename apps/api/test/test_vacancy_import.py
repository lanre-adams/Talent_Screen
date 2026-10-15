"""Vacancy import: the four ECEWS advert sections are captured, with the two
free-text fields (Level of Effort and Minimum Recruitment Standard) kept verbatim."""
from __future__ import annotations

import io

from docx import Document

from app.services.vacancy_import import parse_document, parse_spreadsheet


def _docx(paragraphs: list[tuple[str, str]]) -> bytes:
    doc = Document()
    for text, style in paragraphs:
        if style == "heading":
            doc.add_heading(text, level=1)
        elif style == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_word_label_value_captures_effort_and_standard():
    items = parse_document(_docx([
        ("Data Engineer", "normal"),
        ("Level of Effort: 100%", "normal"),
        ("Description: Build the programme data warehouse", "normal"),
        ("Minimum Recruitment Standard:", "normal"),
        ("Bachelor's degree in a relevant field", "bullet"),
        ("At least 3 years experience", "bullet"),
    ]))
    assert len(items) == 1
    job = items[0].job
    assert job["title"] == "Data Engineer"
    assert job["levelOfEffort"] == "100%"
    assert job["description"] == "Build the programme data warehouse"
    assert "Bachelor's degree" in job["minimumRecruitmentStandard"]
    assert "3 years" in job["minimumRecruitmentStandard"]


def test_word_section_headings_capture_effort_and_standard():
    # The four advert labels appear as Word headings rather than 'Label: value'.
    items = parse_document(_docx([
        ("Job Title", "heading"),
        ("Data Engineer", "normal"),
        ("Level of Effort", "heading"),
        ("100%", "normal"),
        ("Job Description", "heading"),
        ("Build the programme data warehouse", "normal"),
        ("Minimum Recruitment Standard", "heading"),
        ("Bachelor's degree in a relevant field", "bullet"),
        ("At least 3 years experience", "bullet"),
    ]))
    assert len(items) == 1, "section labels must not become new vacancies"
    job = items[0].job
    assert job["title"] == "Data Engineer"
    assert job["levelOfEffort"] == "100%"
    assert job["description"] == "Build the programme data warehouse"
    assert "Bachelor's degree" in job["minimumRecruitmentStandard"]
    assert "3 years" in job["minimumRecruitmentStandard"]


def test_spreadsheet_columns_map_to_effort_and_standard():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Title", "Level of Effort", "Minimum Recruitment Standard", "Department"])
    ws.append(["Data Engineer", "100%", "BSc in a relevant field; 3 years experience", "Strategic Information"])
    buf = io.BytesIO()
    wb.save(buf)

    items = parse_spreadsheet(buf.getvalue())
    assert len(items) == 1
    job = items[0].job
    assert job["title"] == "Data Engineer"
    assert job["levelOfEffort"] == "100%"
    assert job["minimumRecruitmentStandard"] == "BSc in a relevant field; 3 years experience"
