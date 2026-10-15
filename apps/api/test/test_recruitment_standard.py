"""The ECEWS advert "Minimum Recruitment Standard" is turned into scoring criteria."""
from __future__ import annotations

import io

from docx import Document

from app.engine.parser import parse_cv
from app.engine.scorer import score
from app.routers.jobs import normalise_requirements
from app.services.recruitment_standard import (
    normalise_level_of_effort,
    parse_minimum_recruitment_standard,
)
from app.services.vacancy_import import parse_document

W = {"qualifications": 20, "technical": 30, "experience": 30, "certifications": 10, "preferred": 5, "progression": 5}
T = {"shortlist": 85, "consider": 65, "hold": 50}
R = {"enforceMinExperience": True, "enforceMandatory": True, "experienceBasis": "relevant"}

MALARIA_STD = """A first degree in Medicine, Public Health, or a related health field.
A master's degree in Public Health, Epidemiology, or a related field is required.
At least ten (10) years of experience in malaria or infectious disease program management, with a minimum of five (5) years in a technical leadership role.
Strong knowledge of national malaria guidelines and donor-funded malaria programming.
Strong analytical, coordination, and stakeholder engagement skills.
High level of integrity and professional maturity."""

FINANCE_STD = """A first degree in Accounting, Finance, or a related field; a professional accounting qualification (ICAN, ACCA, or equivalent) is required.
At least eight (8) years of progressive experience in financial management, with experience on USG-funded projects preferred.
Strong knowledge of donor financial compliance requirements.
Strong analytical and reporting skills.
High level of integrity and attention to detail."""

DRIVER_STD = """A minimum of a Senior Secondary School Certificate (SSCE).
A valid driver's license and at least five (5) years of professional driving experience.
Good knowledge of road safety regulations and defensive driving practices.
Clean driving record and strong sense of responsibility."""


def test_parses_degree_years_skills_and_industry():
    r = parse_minimum_recruitment_standard(MALARIA_STD)
    assert r["minQualificationLevel"] == 4  # master's required
    assert r["minYears"] == 10
    assert "medicine" in r["fieldsOfStudy"]
    assert "public health" in r["fieldsOfStudy"]
    assert "Malaria" in r["technicalSkills"]
    assert "Public health" in r["technicalSkills"]
    assert "Epidemiology" in r["technicalSkills"]
    assert "Leadership" in r["softSkills"]
    assert "Analytical" in r["softSkills"]
    assert "Stakeholder engagement" in r["softSkills"]


def test_parses_finance_certifications_and_industry():
    r = parse_minimum_recruitment_standard(FINANCE_STD)
    assert r["minQualificationLevel"] == 3
    assert r["minYears"] == 8
    assert r["requiredCertifications"] == ["ICAN", "ACCA"]
    assert "Accounting" in r["technicalSkills"]
    assert "Financial management" in r["technicalSkills"]
    assert "Report writing" in r["technicalSkills"]
    assert r["industryExperience"] == ["NGO"]
    assert "Attention to detail" in r["softSkills"]


def test_parses_secondary_qualification_and_licence():
    r = parse_minimum_recruitment_standard(DRIVER_STD)
    assert r["minQualificationLevel"] == 1
    assert r["minYears"] == 5
    assert r["requiredCertifications"] == ["Driver's license"]


def test_level_of_effort_maps_to_employment_type():
    assert normalise_level_of_effort("Full Time") == "Full-time"
    assert normalise_level_of_effort("Part Time") == "Part-time"
    assert normalise_level_of_effort("Consultancy") == "Consultancy"


def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        if text.startswith("• "):
            doc.add_paragraph(text[2:], style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_ecwes_document_splits_numbered_adverts():
    data = _docx([
        "1.", "Job Title", "Senior Technical Advisor – Malaria",
        "Level of Effort", "Full Time",
        "Job Description",
        "• Provide technical leadership for malaria interventions.",
        "• Guide the scale-up of malaria case management.",
        "Minimum Recruitment Standard",
        "• A first degree in Medicine, Public Health, or a related health field.",
        "• At least ten (10) years of experience in malaria program management.",
        "2.", "Job Title", "State Director",
        "Level of Effort", "Full Time",
        "Job Description", "• Direct ECEWS operations in the assigned state.",
        "Minimum Recruitment Standard",
        "• MB.BS/MD/PHD or similar degree with 10 years of public health experience.",
        "Method of Application",
        "Interested candidates should apply via careers.ecews.org.",
    ])
    items = parse_document(data)
    assert len(items) == 2

    first = items[0]
    assert first.job["title"] == "Senior Technical Advisor – Malaria"
    assert first.job["levelOfEffort"] == "Full Time"
    assert "malaria interventions" in first.job["description"]
    assert "ten (10) years" in first.job["minimumRecruitmentStandard"]

    second = items[1]
    assert second.job["title"] == "State Director"
    # trailing application notes must not leak into the advert
    assert "careers.ecews.org" not in second.job.get("minimumRecruitmentStandard", "")
    assert "careers.ecews.org" not in second.job.get("description", "")


def test_evaluation_is_based_on_parsed_advert():
    """End to end: the parsed standard becomes the scoring criteria for a CV."""
    req = normalise_requirements(parse_minimum_recruitment_standard(MALARIA_STD), require_skill=False)
    cv = parse_cv(
        """Chidinma Okafor
malaria.advisor@example.com

EXPERIENCE
Malaria Technical Advisor — ECEWS | 2016 - present
Led malaria prevention, diagnosis and treatment across supported states.
Senior Malaria Officer — NGO | 2012 - 2016
Supported malaria case management, chemoprevention and vector control.

EDUCATION
MPH Public Health, University of Lagos, 2011
BSc Public Health, University of Ibadan, 2007
"""
    )
    result = score({
        "cv": cv, "jobTitle": "Senior Technical Advisor – Malaria", "req": req,
        "weights": W, "thresholds": T, "rules": R,
    })
    assert result["recommendation"] in ("Shortlist", "Consider")
    assert result["relevantYears"] >= 10
    malaria_check = next(c for c in result["checks"] if c["requirement"] == "Malaria")
    assert malaria_check["met"] is True
