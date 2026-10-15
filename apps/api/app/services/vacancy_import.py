"""Batch vacancy import from an Excel workbook or a Word document.

The same field vocabulary is used for both formats, so a spreadsheet row and a
Word section are parsed into the same shape and then validated by the ordinary
vacancy rules (see `routers/jobs.py: normalise_requirements`).

Excel  — one row per vacancy, the first row that has a "Title" column is the header.
Word   — one Heading 1/2 per vacancy with `Label: value` lines (and bullets),
         a separator line (--- / === / ***) followed by a title, or a table whose
         first row matches the Excel headers.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from ..domain import LEVEL_VALUES

MAX_ROWS = 200


class UnsupportedFile(ValueError):
    """The upload is not a readable .xlsx or .docx."""


# ---------------------------------------------------------------- field vocabulary

def normalise_header(value: str) -> str:
    """'Min Qualification Level' / 'min-qualification_level' → 'minqualificationlevel'."""
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


FIELD_ALIASES: dict[str, str] = {
    # vacancy
    "title": "title", "jobtitle": "title", "vacancy": "title", "vacancytitle": "title",
    "position": "title", "positiontitle": "title", "role": "title", "job": "title",
    "jobname": "title",
    # ECEWS vacancy-advert column: "POSITIONS AVAILABLE"
    "positionsavailable": "title", "positionavailable": "title", "vacanciesavailable": "title",
    "department": "department", "dept": "department", "unit": "department",
    "team": "department",
    "location": "location", "dutystation": "location", "duty": "location",
    "project": "project", "projects": "project", "grant": "project", "programme": "project",
    "numberofpersonsneeded": "numberOfPositions", "numberofpositions": "numberOfPositions",
    "noofpositions": "numberOfPositions", "slots": "numberOfPositions", "headcount": "numberOfPositions",
    "personsneeded": "numberOfPositions", "numberneeded": "numberOfPositions",
    "reportinglines": "reportsTo", "reportsto": "reportsTo", "reportingline": "reportsTo",
    "supervisor": "reportsTo", "linemanager": "reportsTo",
    "employmenttype": "employmentType", "type": "employmentType", "jobtype": "employmentType",
    "contract": "employmentType", "contracttype": "employmentType",
    # ECEWS vacancy-advert section: "LEVEL OF EFFORT"
    "levelofeffort": "levelOfEffort", "effortlevel": "levelOfEffort", "loe": "levelOfEffort",
    "effort": "levelOfEffort",
    "description": "description", "jobdescription": "description", "summary": "description",
    # ECEWS vacancy-advert section: "MINIMUM RECRUITMENT STANDARD"
    "minimumrecruitmentstandard": "minimumRecruitmentStandard",
    "minrecruitmentstandard": "minimumRecruitmentStandard",
    "recruitmentstandard": "minimumRecruitmentStandard",
    "minimumstandard": "minimumRecruitmentStandard",
    "minimumrecruitment": "minimumRecruitmentStandard",
    "closeson": "closesOn", "closingdate": "closesOn", "deadline": "closesOn",
    "applicationsclose": "closesOn", "closing": "closesOn",
    # requirements
    "minqualificationlevel": "minQualificationLevel", "minimumqualificationlevel": "minQualificationLevel",
    "minimumqualification": "minQualificationLevel", "minqualification": "minQualificationLevel",
    "qualificationlevel": "minQualificationLevel", "qualification": "minQualificationLevel",
    "minlevel": "minQualificationLevel", "education": "minQualificationLevel",
    "fieldsofstudy": "fieldsOfStudy", "fieldofstudy": "fieldsOfStudy", "fields": "fieldsOfStudy",
    "discipline": "fieldsOfStudy",
    "requiredqualifications": "requiredQualifications", "requiredqualification": "requiredQualifications",
    "qualifications": "requiredQualifications",
    "preferredqualifications": "preferredQualifications", "preferredqualification": "preferredQualifications",
    "requiredcertifications": "requiredCertifications", "requiredcertification": "requiredCertifications",
    "certifications": "requiredCertifications", "certification": "requiredCertifications",
    "technicalskills": "technicalSkills", "technicalskill": "technicalSkills",
    "requiredskills": "technicalSkills", "skills": "technicalSkills", "technical": "technicalSkills",
    "softskills": "softSkills", "softskill": "softSkills",
    "preferredskills": "preferredSkills", "preferredskill": "preferredSkills",
    "desirableskills": "preferredSkills",
    "minyears": "minYears", "minimumyears": "minYears", "yearsofexperience": "minYears",
    "minyearsofexperience": "minYears", "minexperience": "minYears", "experience": "minYears",
    "minyrs": "minYears", "years": "minYears",
    "preferredyears": "preferredYears", "preferredyearsofexperience": "preferredYears",
    "preferredexperience": "preferredYears",
    "industryexperience": "industryExperience", "industry": "industryExperience",
    "sector": "industryExperience", "sectorexperience": "industryExperience",
    "locationrequirements": "locationRequirements", "locationrequirement": "locationRequirements",
    "mandatory": "mandatory", "mandatoryrequirements": "mandatory", "essential": "mandatory",
    "nicetohave": "niceToHave", "nicetohaves": "niceToHave", "desirable": "niceToHave",
    "optional": "niceToHave",
}

JOB_FIELDS = ("title", "department", "location", "employmentType", "description", "closesOn",
              "project", "reportsTo", "levelOfEffort", "minimumRecruitmentStandard")
JOB_NUMBER_FIELDS = ("numberOfPositions",)
REQUIREMENT_LIST_FIELDS = (
    "fieldsOfStudy", "requiredQualifications", "preferredQualifications", "requiredCertifications",
    "technicalSkills", "softSkills", "preferredSkills", "industryExperience", "mandatory", "niceToHave",
)
REQUIREMENT_TEXT_FIELDS = ("locationRequirements",)
REQUIREMENT_NUMBER_FIELDS = ("minYears", "preferredYears")

# The four labels every ECEWS vacancy advert carries. In a Word document they act
# as section markers: content between them is captured verbatim into that field.
SECTION_FIELDS: dict[str, str] = {
    "jobtitle": "title",
    "levelofeffort": "levelOfEffort",
    "jobdescription": "description",
    "minimumrecruitmentstandard": "minimumRecruitmentStandard",
}

# Front/back matter that ends the advert content in an ECEWS document.
_BACK_MATTER = {
    "note", "methodofapplication", "applicationmust", "candidatesareadvised",
    "eligiblefemale", "onlyshortlisted", "interestedandqualified",
}

_ORDERING = re.compile(r"^\d+\s*\.?$")


def _clean_line(text: str) -> str:
    """Normalise a paragraph: NBSPs → spaces, strip list glyphs, collapse whitespace."""
    t = text.replace("\xa0", " ").strip()
    t = re.sub(r"^[\u00b7\u2022\u25aa\u2023\u2013\u2014*\-]\s*", "", t).strip()
    return re.sub(r"\s+", " ", t)


def _parse_ecwes(paragraphs: list[tuple[str, str]]) -> list[Item]:
    """Parse the four-section ECEWS advert format.

    Each vacancy is a "Job Title" label followed by its title, then "Level of
    Effort", "Job Description" and "Minimum Recruitment Standard" sections whose
    lines accumulate verbatim into the corresponding field. Numbered list markers
    ("1.", "33.") and trailing application notes are ignored.
    """
    items: list[Item] = []
    current: Item | None = None
    open_field: str | None = None

    for text, _style in paragraphs:
        t = _clean_line(text)
        if not t or _ORDERING.match(t):
            continue
        # Colon form ("Level of Effort: 100%") as a convenience.
        key, value = _split_label(t) if ":" in t else (None, "")
        if key:
            if key == "title":
                current = Item(f"Section {len(items) + 1}")
                items.append(current)
                open_field = "title"
                if value:
                    current.job["title"] = value
                    open_field = None
            else:
                open_field = key
                if value and current is not None:
                    current.set(key, value)
            continue
        section = SECTION_FIELDS.get(normalise_header(t))
        if section:
            if section == "title":
                current = Item(f"Section {len(items) + 1}")
                items.append(current)
                open_field = "title"
            else:
                open_field = section
            continue
        if normalise_header(t) in _BACK_MATTER:
            break
        if current is None:
            continue
        if open_field == "title":
            current.job["title"] = t
            open_field = None
        elif open_field:
            previous = current.job.get(open_field) or ""
            current.set(open_field, f"{previous}\n{t}" if previous else t)
    return items

# Human-readable header for each canonical key, in template order.
TEMPLATE_COLUMNS: list[tuple[str, str, str]] = [
    ("title", "Title", "Data Engineer"),
    ("department", "Department", "Strategic Information"),
    ("location", "Location", "Abuja, Nigeria"),
    ("project", "Project", "BRIDGE"),
    ("numberOfPositions", "Number of Positions", "3"),
    ("reportsTo", "Reports To", "State Director"),
    ("employmentType", "Employment Type", "Full-time"),
    ("levelOfEffort", "Level of Effort", "100%"),
    ("closesOn", "Closes On", "2026-12-31"),
    ("description", "Description", "Build and maintain the programme data warehouse."),
    ("minimumRecruitmentStandard", "Minimum Recruitment Standard", "BSc in a relevant field; at least 3 years experience"),
    ("minQualificationLevel", "Min Qualification Level", "3"),
    ("fieldsOfStudy", "Fields of Study", "Computer Science; Statistics"),
    ("requiredQualifications", "Required Qualifications", "BSc Computer Science or related"),
    ("preferredQualifications", "Preferred Qualifications", "MSc in a data discipline"),
    ("requiredCertifications", "Required Certifications", "Azure Administrator (AZ-104)"),
    ("technicalSkills", "Technical Skills", "SQL; Python; DHIS2; Power BI"),
    ("softSkills", "Soft Skills", "Communication; Teamwork"),
    ("preferredSkills", "Preferred Skills", "KoboToolbox"),
    ("minYears", "Min Years", "3"),
    ("preferredYears", "Preferred Years", "5"),
    ("industryExperience", "Industry Experience", "NGO; Public health"),
    ("locationRequirements", "Location Requirements", "Must be based in Abuja"),
    ("mandatory", "Mandatory", "Willing to travel to field sites"),
    ("niceToHave", "Nice-to-have", "French language"),
]

# Ordered so that longer, more specific labels win ("postgraduate diploma" beats "diploma").
_LEVEL_KEYWORDS: list[tuple[tuple[str, ...], float]] = [
    (("phd", "doctorate", "doctoral"), 5),
    (("postgraduate diploma", "pgd", "post graduate diploma"), 3.5),
    (("master", "msc", "mba", "meng", "masters"), 4),
    (("bachelor", "bsc", "hnd", "beng", "btech", "ba ", "degree", "b.a"), 3),
    (("diploma", "ond", "nd ", "national diploma"), 2),
    (("secondary", "ssce", "waec", "olevel", "o level", "high school", "ssc"), 1),
]


def parse_qualification(value: Any) -> float | None:
    """'3', 3, 'Bachelor's / HND' → 3.0; unknown → None."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if float(value) in LEVEL_VALUES else None
    text = str(value).strip().lower()
    if not text:
        return None
    try:
        number = float(text)
        return number if number in LEVEL_VALUES else None
    except ValueError:
        pass
    for keywords, level in _LEVEL_KEYWORDS:
        if any(k in text for k in keywords):
            return level
    return None


def parse_number(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    match = re.search(r"-?\d+(\.\d+)?", text)
    return float(match.group(0)) if match else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class Item:
    """One parsed vacancy, with any problems found while reading it."""

    __slots__ = ("label", "job", "requirements", "warnings")

    def __init__(self, label: str) -> None:
        self.label = label
        self.job: dict[str, Any] = {}
        self.requirements: dict[str, Any] = {}
        self.warnings: list[str] = []

    def set(self, key: str, raw: Any) -> None:
        if key in JOB_FIELDS:
            text = _cell_text(raw)
            if text:
                self.job[key] = text
            return
        if key in JOB_NUMBER_FIELDS:
            number = parse_number(raw)
            if number is None:
                if _cell_text(raw):
                    self.warnings.append(f"Could not read {key} {raw!r}; it was ignored")
            else:
                self.job[key] = int(number)
            return
        if key == "minQualificationLevel":
            level = parse_qualification(raw)
            if level is None:
                if _cell_text(raw):
                    self.warnings.append(
                        f"Could not read qualification {raw!r}; the default level was used"
                    )
            else:
                self.requirements[key] = level
            return
        if key in REQUIREMENT_NUMBER_FIELDS:
            number = parse_number(raw)
            if number is None:
                if _cell_text(raw):
                    self.warnings.append(f"Could not read {key} {raw!r}; it was ignored")
            else:
                self.requirements[key] = number
            return
        text = _cell_text(raw)
        if text:
            self.requirements[key] = text

    def json(self) -> dict[str, Any]:
        return {"label": self.label, "job": self.job, "requirements": self.requirements,
                "warnings": list(self.warnings)}


def _items_from_rows(headers: list[str], rows: list[tuple], start: int) -> list[Item]:
    """Shared by the Excel reader and the Word-table reader."""
    columns = [FIELD_ALIASES.get(normalise_header(h)) for h in headers]
    items: list[Item] = []
    for offset, row in enumerate(rows):
        if len(items) >= MAX_ROWS:
            break
        item = Item(f"Row {start + offset}")
        for index, key in enumerate(columns):
            if key and index < len(row):
                item.set(key, row[index])
        if any(item.job.values()) or item.requirements:
            items.append(item)
    return items


# ------------------------------------------------------------------------- Excel


def parse_spreadsheet(data: bytes) -> list[Item]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise UnsupportedFile(f"Could not read the spreadsheet ({e})")
    sheet = workbook.worksheets[0]
    rows = [tuple(r) for r in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        raise UnsupportedFile("The spreadsheet is empty")

    header_index = -1
    header: list[str] = []
    for i, row in enumerate(rows):
        candidates = [normalise_header(_cell_text(c)) for c in row]
        if any(FIELD_ALIASES.get(c) == "title" for c in candidates):
            header_index = i
            header = [str(_cell_text(c)) for c in row]
            break
    if header_index < 0:
        raise UnsupportedFile('No header row found — the sheet needs a "Title" column')
    unknown = [h for h in header if h and FIELD_ALIASES.get(normalise_header(h)) is None]
    items = _items_from_rows(header, rows[header_index + 1:], header_index + 2)
    if unknown:
        # Attach to the first item so the caller can report it once.
        note = f"Ignored unrecognised column(s): {', '.join(unknown)}"
        if items:
            items[0].warnings.append(note)
        else:
            raise UnsupportedFile(f"{note}. No data rows were found.")
    return items


# -------------------------------------------------------------------------- Word


def _split_label(text: str) -> tuple[str | None, str]:
    """'Technical Skills: SQL' → ('technicalSkills', 'SQL')."""
    line = text.strip().lstrip("-•*–\t ").strip()
    if ":" not in line:
        return None, ""
    head, _, tail = line.partition(":")
    return FIELD_ALIASES.get(normalise_header(head)), tail.strip()


def _is_separator(text: str) -> bool:
    return bool(re.fullmatch(r"[-=*_]{3,}", text.strip()))


def parse_document(data: bytes) -> list[Item]:
    try:
        from docx import Document  # python-docx

        document = Document(io.BytesIO(data))
    except Exception as e:
        raise UnsupportedFile(f"Could not read the document ({e})")

    paragraphs = [
        (p.text.strip(), (getattr(p.style, "name", "") or "").lower())
        for p in document.paragraphs
    ]
    has_heading = any(style.startswith("heading") for text, style in paragraphs if text)

    # A "Minimum Recruitment Standard" label (no colon) marks the ECEWS advert
    # format. Route it to the dedicated parser before the generic heading/label logic.
    if any(
        t and ":" not in t and normalise_header(t) == "minimumrecruitmentstandard"
        for t, _ in paragraphs
    ):
        return _parse_ecwes(paragraphs)

    items: list[Item] = []
    current: Item | None = None
    open_field: str | None = None
    expect_title = False

    def start() -> Item:
        item = Item(f"Section {len(items) + 1}")
        items.append(item)
        return item

    if has_heading:
        # Heading-driven: each Heading starts a vacancy. Anything before the first
        # heading is front matter (document title, notes), so it is ignored.
        for text, style in paragraphs:
            if not text:
                continue
            # The four ECEWS advert labels are section markers, not new vacancies.
            section = SECTION_FIELDS.get(normalise_header(text))
            if section:
                if section == "title":
                    current = start()
                    open_field = None
                else:
                    open_field = section
                continue
            if style.startswith("heading"):
                current = start()
                current.job["title"] = text
                open_field = None
                continue
            if current is None:
                continue
            if _is_separator(text):
                continue
            key, value = _split_label(text)
            if key:
                open_field = key
                if value:
                    current.set(key, value)
                continue
            if open_field:
                # A bullet or wrapped line belonging to the label above.
                previous = current.requirements.get(open_field) or current.job.get(open_field) or ""
                current.set(open_field, f"{previous}\n{text}" if previous else text)
                continue
            if "title" not in current.job:
                current.job["title"] = text
            else:
                current.set("description", f"{current.job.get('description', '')}\n{text}".strip())
    else:
        # No headings: '---' separators (title on the next line), otherwise one vacancy.
        for text, _style in paragraphs:
            if not text:
                continue
            if _is_separator(text):
                expect_title = True
                continue
            # The four ECEWS advert labels are section markers, not titles or text.
            section = SECTION_FIELDS.get(normalise_header(text))
            if section:
                if section == "title":
                    current = start()
                    open_field = None
                else:
                    open_field = section
                continue
            if expect_title:
                current = start()
                current.job["title"] = text
                expect_title = False
                open_field = None
                continue
            if current is None:
                current = start()
            key, value = _split_label(text)
            if key:
                open_field = key
                if value:
                    current.set(key, value)
                continue
            if open_field:
                previous = current.requirements.get(open_field) or current.job.get(open_field) or ""
                current.set(open_field, f"{previous}\n{text}" if previous else text)
                continue
            if "title" not in current.job:
                current.job["title"] = text
            else:
                current.set("description", f"{current.job.get('description', '')}\n{text}".strip())

    # A Word document may instead carry the same table as the spreadsheet.
    for table in document.tables:
        rows = [tuple(_cell_text(c.text) for c in r.cells) for r in table.rows]
        if not rows:
            continue
        header = list(rows[0])
        if not any(FIELD_ALIASES.get(normalise_header(h)) == "title" for h in header):
            continue
        table_items = _items_from_rows(header, rows[1:], 2)
        for offset, item in enumerate(table_items):
            item.label = f"Table {len(items) + offset + 1}"
        items.extend(table_items)

    if not items:
        raise UnsupportedFile(
            "No vacancies were found. Use one Heading per vacancy with 'Label: value' lines, "
            "or download the Word template."
        )
    return items


# ----------------------------------------------------------------------- facade


def parse_upload(filename: str, data: bytes) -> list[Item]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return parse_spreadsheet(data)
    if name.endswith(".docx"):
        return parse_document(data)
    if name.endswith(".xls") or name.endswith(".doc"):
        raise UnsupportedFile("The old .xls/.doc formats are not supported — save as .xlsx or .docx")
    raise UnsupportedFile("Upload an .xlsx spreadsheet or a .docx document")


# -------------------------------------------------------------------- templates


def build_template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vacancies"
    sheet.append([heading for _, heading, _ in TEMPLATE_COLUMNS])
    sheet.append([example for _, _, example in TEMPLATE_COLUMNS])
    sheet.append(["" for _ in TEMPLATE_COLUMNS])

    for index in range(1, len(TEMPLATE_COLUMNS) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", start_color="FF316C4E", end_color="FF316C4E")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = max(16, len(TEMPLATE_COLUMNS[index - 1][1]) + 4)

    notes = workbook.create_sheet("Instructions")
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 96
    notes.append(["Column", "What to put in it"])
    notes["A1"].font = Font(bold=True)
    notes["B1"].font = Font(bold=True)
    notes.append(["Title", "Required. One row per vacancy — a row without a title is skipped."])
    notes.append(["Department", "Must match an existing department name. Left blank if it does not match."])
    notes.append(["Project / Number of Positions / Reports To", "Optional. Stored on the vacancy and shown on its page (ECEWS vacancy-advert columns)."])
    notes.append(["Employment Type", "Full-time, Part-time, Contract, Consultancy, Internship or Temporary. Defaults to Full-time."])
    notes.append(["Level of Effort", "Free text, e.g. 100% (ECEWS vacancy-advert column)."])
    notes.append(["Closes On", "A date, e.g. 2026-12-31. Optional."])
    notes.append(["Description", "Free text. Line breaks are preserved."])
    notes.append(["Minimum Recruitment Standard", "Free text describing the minimum criteria, kept verbatim (not parsed into scoring fields)."])
    notes.append(["Min Qualification Level", "A number 1-5 (3.5 = Postgraduate Diploma), or a name such as Bachelor's or Master's."])
    notes.append(["Min Years / Preferred Years", "Numbers. Preferred must be at least the minimum."])
    notes.append(["Skill / list columns", "Separate items with a semicolon, a comma or a new line inside the cell."])
    notes.append(["Technical Skills", "Required: at least one technical skill per vacancy."])
    notes.append(["", ""])
    notes.append(["Unknown columns", "Ignored, and reported back after the upload."])
    notes.append(["Row limit", f"At most {MAX_ROWS} vacancies per upload."])
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_template_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Vacancy batch upload — template", level=0)
    document.add_paragraph(
        "Each Heading 1 starts a new vacancy. Under it, write one 'Label: value' line per field; "
        "a label with nothing after the colon collects the bullet lines that follow it. "
        "If your document uses no headings at all, a line of --- starts a new vacancy instead and the "
        "next line becomes its title. Delete the examples below before adding your own."
    )

    document.add_heading("Data Engineer", level=1)
    document.add_paragraph("Department: Strategic Information")
    document.add_paragraph("Location: Abuja, Nigeria")
    document.add_paragraph("Project: BRIDGE")
    document.add_paragraph("Number of Positions: 3")
    document.add_paragraph("Reports To: State Director")
    document.add_paragraph("Employment Type: Full-time")
    document.add_paragraph("Level of Effort: 100%")
    document.add_paragraph("Closes On: 2026-12-31")
    document.add_paragraph("Min Qualification Level: Bachelor's")
    document.add_paragraph("Fields of Study: Computer Science; Statistics")
    document.add_paragraph("Min Years: 3")
    document.add_paragraph("Preferred Years: 5")
    document.add_paragraph("Industry Experience: NGO; Public health")
    document.add_paragraph("Technical Skills:")
    for skill in ("SQL", "Python", "DHIS2", "Power BI"):
        document.add_paragraph(skill, style="List Bullet")
    document.add_paragraph("Soft Skills:")
    for skill in ("Communication", "Teamwork"):
        document.add_paragraph(skill, style="List Bullet")
    document.add_paragraph("Mandatory:")
    document.add_paragraph("Willing to travel to field sites", style="List Bullet")
    document.add_paragraph("Description: Build and maintain the programme data warehouse.")
    document.add_paragraph("Minimum Recruitment Standard:")
    document.add_paragraph("Bachelor's degree in a relevant field", style="List Bullet")
    document.add_paragraph("At least 3 years experience", style="List Bullet")

    document.add_heading("Programme Officer", level=1)
    document.add_paragraph("Department: Programmes")
    document.add_paragraph("Technical Skills: Report writing; Stakeholder management")
    document.add_paragraph("Min Years: 2")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
